# -*- coding: utf-8 -*-


class DocumentReader:
    def read(self, file_path):
        return read_document_file(file_path)


def read_txt_file(file_path):
    """Đọc nội dung file .txt bằng Python thuần."""
    encodings = ["utf-8", "utf-8-sig", "cp1258", "latin-1"]
    last_error = None
    for encoding in encodings:
        try:
            with open(file_path, "r", encoding=encoding) as file:
                return file.read()
        except UnicodeDecodeError as error:
            last_error = error
    raise ValueError(f"Không đọc được file TXT. Lỗi mã hóa: {last_error}")


def read_docx_file(file_path):
    """Đọc nội dung file Word .docx."""
    try:
        from docx import Document
    except ImportError as error:
        raise ImportError("Thiếu thư viện python-docx. Hãy chạy: pip install python-docx") from error

    document = Document(file_path)
    text_parts = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            row_text = " ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)

    return "\n".join(text_parts)


def read_pdf_file(file_path):
    """Đọc nội dung file PDF.

    Lưu ý: PDF scan dạng ảnh cần OCR nên hàm này có thể không trích xuất được chữ.
    """
    try:
        import PyPDF2
    except ImportError as error:
        raise ImportError("Thiếu thư viện PyPDF2. Hãy chạy: pip install PyPDF2") from error

    text_parts = []
    with open(file_path, "rb") as file:
        reader = PyPDF2.PdfReader(file)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text and page_text.strip():
                text_parts.append(page_text.strip())

    return "\n".join(text_parts)


def read_document_file(file_path):
    """Đọc văn bản từ .txt, .docx hoặc .pdf để đưa vào hệ thống phân loại."""
    path = str(file_path)
    lower_path = path.lower()

    if lower_path.endswith(".txt"):
        return read_txt_file(path)
    if lower_path.endswith(".docx"):
        return read_docx_file(path)
    if lower_path.endswith(".pdf"):
        return read_pdf_file(path)

    raise ValueError("Chỉ hỗ trợ import file .txt, .docx hoặc .pdf")
